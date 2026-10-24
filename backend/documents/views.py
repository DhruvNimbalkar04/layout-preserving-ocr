import os
import json
import threading

from django.db import close_old_connections
from django.shortcuts import get_object_or_404
from django.http import HttpResponse, Http404
from django.utils import timezone
from rest_framework.generics import ListCreateAPIView, RetrieveDestroyAPIView
from rest_framework.views import APIView
from rest_framework.response import Response
from rest_framework import status
from .models import Documents, ExtractedText
from .serializers import DocumentSerializer
from .pymupdf_extractor import (
    LAYOUT_VERSION,
    extract_document_layout,
    render_page_image,
    generate_reconstructed_pdf,
    generate_layout_html,
)


# ---------------------------------------------------------------------------
# Background extraction jobs
#
# OCR can take several minutes on scanned documents, which is far longer than a
# browser/proxy will wait. Extraction therefore runs in a worker thread and the
# API reports progress; at most one job per document may be in flight.
# ---------------------------------------------------------------------------
_jobs = {}
_jobs_lock = threading.Lock()


def _flatten_layout_text(layout):
    lines = []
    for page in layout.get("pages", []):
        for el in page.get("elements", []):
            if el.get("type") != "image" and el.get("text"):
                lines.append(el["text"])
    return "\n".join(lines)


def _job_snapshot(document_id):
    with _jobs_lock:
        job = _jobs.get(document_id)
        return dict(job) if job else None


def _update_job(document_id, **fields):
    with _jobs_lock:
        job = _jobs.setdefault(
            document_id,
            {"state": "processing", "page": 0, "total": 0, "error": ""},
        )
        job.update(fields)


def _job_payload(job):
    job = job or {}
    return {
        "status": job.get("state", "processing"),
        "page": job.get("page", 0),
        "total": job.get("total", 0),
        "error": job.get("error", ""),
    }


def _run_extraction(document_id):
    close_old_connections()
    try:
        document = Documents.objects.get(pk=document_id)
        use_ocr = getattr(document, "document_kind", "normal") == "scanned"

        def on_progress(page_number, total_pages):
            _update_job(document_id, page=page_number, total=total_pages)

        layout = extract_document_layout(
            document.file.path, use_ocr=use_ocr, on_progress=on_progress
        )

        extracted_obj, _ = ExtractedText.objects.get_or_create(document=document)
        extracted_obj.layout_data = layout
        extracted_obj.text = _flatten_layout_text(layout)
        extracted_obj.save()

        document.status = "completed"
        document.save(update_fields=["status", "updated_at"])

        pages = layout.get("pages", [])
        _update_job(
            document_id,
            state="completed",
            page=len(pages),
            total=len(pages),
            error="",
        )
    except Exception as err:
        print(f"Extraction failed for document {document_id}: {err}")
        try:
            Documents.objects.filter(pk=document_id).update(
                status="failed", updated_at=timezone.now()
            )
        except Exception:
            pass
        _update_job(document_id, state="failed", error=str(err))
    finally:
        close_old_connections()


def _start_extraction(document_id, document=None):
    """Start one extraction worker per document. Returns False if one is running."""
    with _jobs_lock:
        job = _jobs.get(document_id)
        if job and job.get("state") == "processing":
            return False
        _jobs[document_id] = {"state": "processing", "page": 0, "total": 0, "error": ""}

    if document is not None:
        document.status = "processing"
        document.save(update_fields=["status", "updated_at"])
    else:
        Documents.objects.filter(pk=document_id).update(
            status="processing", updated_at=timezone.now()
        )

    threading.Thread(
        target=_run_extraction,
        args=(document_id,),
        name=f"ocr-document-{document_id}",
        daemon=True,
    ).start()
    return True


def _has_cached_layout(extracted_obj):
    layout = extracted_obj.layout_data
    return bool(
        layout
        and isinstance(layout, dict)
        and layout.get("pages")
        and layout.get("layout_version") == LAYOUT_VERSION
    )


class DocumentView(ListCreateAPIView):
    queryset = Documents.objects.all().order_by('-created_at')
    serializer_class = DocumentSerializer

    def perform_create(self, serializer):
        document = serializer.save()
        # Kick off extraction right away so the viewer usually finds it ready.
        _start_extraction(document.id, document=document)


class DocumentDetailView(RetrieveDestroyAPIView):
    queryset = Documents.objects.all()
    serializer_class = DocumentSerializer


class DocumentPageView(APIView):
    def get(self, request, pk, page_num):
        doc = get_object_or_404(Documents, pk=pk)
        if not doc.file or not doc.file.storage.exists(doc.file.name):
            raise Http404("Document file not found")

        try:
            img_bytes, content_type = render_page_image(doc.file.path, page_num)
            if img_bytes is None:
                raise Http404("Page not found")
            return HttpResponse(img_bytes, content_type=content_type)
        except Http404:
            raise
        except Exception as e:
            raise Http404(f"Error rendering page: {str(e)}")


class DocumentLayoutView(APIView):
    def get(self, request, document_id):
        document = get_object_or_404(Documents, pk=document_id)
        extracted_obj, _ = ExtractedText.objects.get_or_create(document=document)

        force_refresh = request.query_params.get("refresh", "false").lower() == "true"
        has_cached = _has_cached_layout(extracted_obj)

        if has_cached and not force_refresh:
            return Response(extracted_obj.layout_data)

        job = _job_snapshot(document_id)
        if job and job["state"] == "processing":
            return Response(_job_payload(job), status=status.HTTP_202_ACCEPTED)

        if job and job["state"] == "failed" and not force_refresh:
            return Response(_job_payload(job))

        _start_extraction(document_id, document=document)
        return Response(
            _job_payload(_job_snapshot(document_id)),
            status=status.HTTP_202_ACCEPTED,
        )

    def patch(self, request, document_id):
        document = get_object_or_404(Documents, pk=document_id)
        extracted_obj, _ = ExtractedText.objects.get_or_create(document=document)

        layout_data = request.data.get("layout_data") or request.data
        if not isinstance(layout_data, dict) or "pages" not in layout_data:
            return Response(
                {"error": "Invalid layout data format. Must include 'pages' key."},
                status=status.HTTP_400_BAD_REQUEST
            )

        extracted_obj.layout_data = layout_data
        extracted_obj.text = _flatten_layout_text(layout_data)
        extracted_obj.save()

        return Response(extracted_obj.layout_data, status=status.HTTP_200_OK)


class DocumentStatusView(APIView):
    def get(self, request, document_id):
        document = get_object_or_404(Documents, pk=document_id)
        extracted_obj = ExtractedText.objects.filter(document=document).first()
        has_layout = bool(extracted_obj and _has_cached_layout(extracted_obj))

        job = _job_snapshot(document_id)
        if not job and not has_layout and document.status in ("pending", "processing"):
            # No in-memory job (e.g. the dev server restarted mid-extraction):
            # start it again so polling heals itself.
            _start_extraction(document_id, document=document)
            job = _job_snapshot(document_id)

        payload = _job_payload(job) if job else {
            "status": document.status,
            "page": 0,
            "total": 0,
            "error": "",
        }
        payload["has_layout"] = has_layout
        return Response(payload)



class DocumentExportView(APIView):
    def get(self, request, document_id):
        document = get_object_or_404(Documents, pk=document_id)
        extracted_obj, _ = ExtractedText.objects.get_or_create(document=document)
        export_format = request.query_params.get("export_format", "pdf").lower()
        base_name = os.path.splitext(document.file_name)[0]

        layout = extracted_obj.layout_data or {}
        pages = layout.get("pages", [])

        if export_format == "json":
            response = HttpResponse(json.dumps(layout, indent=2), content_type="application/json")
            response["Content-Disposition"] = f'attachment; filename="{base_name}_layout.json"'
            return response

        elif export_format == "txt":
            text_content = extracted_obj.text or ""
            if not text_content and pages:
                all_text = []
                for p in pages:
                    p_text = "\n".join(
                        el["text"] for el in p.get("elements", [])
                        if el.get("type") != "image" and el.get("text")
                    )
                    if p_text:
                        all_text.append(f"--- Page {p.get('page_number', 1)} ---\n" + p_text)
                text_content = "\n\n".join(all_text)
            response = HttpResponse(text_content, content_type="text/plain; charset=utf-8")
            response["Content-Disposition"] = f'attachment; filename="{base_name}_extracted.txt"'
            return response

        elif export_format == "html":
            html_content = generate_layout_html(layout, document.file_name)
            response = HttpResponse(html_content, content_type="text/html; charset=utf-8")
            response["Content-Disposition"] = f'attachment; filename="{base_name}.html"'
            return response

        elif export_format == "pdf":
            pdf_bytes = generate_reconstructed_pdf(document, layout)
            response = HttpResponse(pdf_bytes, content_type="application/pdf")
            response["Content-Disposition"] = f'attachment; filename="{base_name}_edited.pdf"'
            return response

        elif export_format in ("docx", "doc"):
            import io
            import base64
            import docx
            from docx.shared import Pt, Inches
            from docx.enum.text import WD_ALIGN_PARAGRAPH

            doc_file = docx.Document()
            for sec in doc_file.sections:
                sec.top_margin = Inches(1)
                sec.bottom_margin = Inches(1)
                sec.left_margin = Inches(1)
                sec.right_margin = Inches(1)

            for page_idx, page_data in enumerate(pages):
                if page_idx > 0:
                    doc_file.add_page_break()

                pw = page_data.get("width", 595.2)
                pc = pw / 2

                for el in page_data.get("elements", []):
                    if el.get("type") == "image":
                        src = el.get("src", "")
                        if src.startswith("data:") and not el.get("isBackground"):
                            try:
                                _, b64_data = src.split(",", 1)
                                img_bytes = base64.b64decode(b64_data)
                                img_w = min(3.5, max(1.0, el.get("bbox", {}).get("width", 150) / 72.0))
                                p = doc_file.add_paragraph()
                                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                                run = p.add_run()
                                run.add_picture(io.BytesIO(img_bytes), width=Inches(img_w))
                            except Exception as err:
                                print("Error adding image to docx:", err)
                    else:
                        text = el.get("text", "").strip()
                        if not text:
                            continue

                        bbox = el.get("bbox", {})
                        ec = bbox.get("x", 0) + bbox.get("width", 0) / 2
                        is_center = abs(ec - pc) < 45

                        p = doc_file.add_paragraph()
                        if is_center:
                            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        elif bbox.get("x", 0) > pw * 0.6:
                            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                        else:
                            p.alignment = WD_ALIGN_PARAGRAPH.LEFT

                        style = el.get("style", {})
                        run = p.add_run(text)
                        run.font.name = "Times New Roman"
                        fontsize = style.get("fontSize", 12)
                        run.font.size = Pt(max(9, min(26, fontsize)))
                        run.font.bold = bool(style.get("bold"))
                        run.font.italic = bool(style.get("italic"))

            docx_buf = io.BytesIO()
            doc_file.save(docx_buf)
            response = HttpResponse(
                docx_buf.getvalue(),
                content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            )
            response["Content-Disposition"] = f'attachment; filename="{base_name}.docx"'
            return response

        return Response(
            {"error": f"Unsupported format: {export_format}"},
            status=status.HTTP_400_BAD_REQUEST
        )